"""Verify rebuilt Q1 and prove Q2/Q3 files, sheets and responses were preserved."""
from pathlib import Path
import json,sys,hashlib
try:import openpyxl
except ImportError:
    sys.path.append(str(Path.home()/'.cache/codex-runtimes/codex-primary-runtime/dependencies/python/Lib/site-packages'))
    import openpyxl
import numpy as np
import q1_beam_column as q
ROOT=Path(__file__).resolve().parent.parent;BASE=ROOT/'checkpoints/before_Q1_rebuild'
hashes=json.loads((BASE/'hashes.json').read_text())
for name,h in hashes.items():assert hashlib.sha256((ROOT/name).read_bytes()).hexdigest()==h,name
old=openpyxl.load_workbook(BASE/'CE587_HW1_Solver.xlsx');new=openpyxl.load_workbook(ROOT/'CE587_HW1_Solver.xlsx')
olddata=openpyxl.load_workbook(BASE/'CE587_HW1_Solver.xlsx',data_only=True)
data=openpyxl.load_workbook(ROOT/'CE587_HW1_Solver.xlsx',data_only=True)
preserved=[n for n in old.sheetnames if n.startswith(('Q2','Q3'))];count=0
for name in preserved:
    a=new[name];b=old[name]
    assert len(a._images)==len(b._images),name
    assert a.print_area==b.print_area,(name,'print area')
    for row in b:
        for c in row:
            v=a[c.coordinate].value;ref=c.value
            if hasattr(ref,'text'):assert v.text==ref.text and v.ref==ref.ref,(name,c.coordinate)
            elif isinstance(ref,(int,float)):assert np.isclose(v,ref,rtol=1e-14,atol=1e-14),(name,c.coordinate)
            else:assert v==ref,(name,c.coordinate)
            assert a[c.coordinate].number_format==c.number_format,(name,c.coordinate,'format')
            if ref is not None:count+=1
    for row in olddata[name]:
        for c in row:
            v=data[name][c.coordinate].value;ref=c.value
            if isinstance(ref,(int,float)):assert np.isclose(v,ref,rtol=1e-13,atol=1e-12),(name,c.coordinate,'cached')
            else:assert v==ref,(name,c.coordinate,'cached')
for name,d in old.defined_names.items():
    if name.startswith(('Q2_','Q3_')):assert new.defined_names[name].attr_text==d.attr_text,name
s=data['Q1_BeamColumn'];x=q.cm_to_m([s['B7'].value,s['B8'].value]);ind=q.independent_check(x)
for cell,name in [('B4','mass_kg'),('B20','area_m2'),('B21','inertia_m4'),('B22','self_weight_N_m'),('B23','stress_Pa'),('B24','deflection_m'),('B25','radius_thickness_ratio')]:
    assert np.isclose(s[cell].value,ind[name],rtol=1e-12,atol=1e-14),(cell,name)
assert q.feasible(x)
assert abs(q.mass(x)-q.analytic()['mass_kg'])<1e-6
errors=[(sh.title,c.coordinate,c.value) for sh in data for row in sh for c in row if c.data_type=='e']
assert not errors,errors
assert not any('#REF!' in str(d.attr_text) for d in new.defined_names.values())
assert s['J1'].value is not None
out=dict(q2_q3_source_result_and_figure_hashes_unchanged=True,preserved_existing_sheets=preserved,
    preserved_populated_cells=count,q2_q3_formulas_values_formats_cached_responses_named_inputs_and_image_counts='passed',
    formula_errors=errors,q1_excel_python_response_parity=True,q1_excel_x_cm=q.m_to_cm(x).tolist(),
    q1_excel_independent=ind,q1_excel_g=q.constraints(x).tolist(),q1_excel_activity=q.activity(x),
    q1_excel_mass_error_kg=float(q.mass(x)-4.68),native_q1_solver_model_present=True)
(ROOT/'verification/q1_project.json').write_text(json.dumps(out,indent=2));print(json.dumps(out,indent=2))
